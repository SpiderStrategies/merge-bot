#!/usr/bin/env python3
"""Generates the 'How the Spider Merge Bot Works' presentation (.pptx).

Conceptual intro for new developers. Run: python3 docs/build_presentation.py
Produces: docs/spider-merge-bot.pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- palette ---------------------------------------------------------------
INK = RGBColor(0x1A, 0x1A, 0x2E)        # near-black for body text
ACCENT = RGBColor(0x2E, 0x6F, 0xF2)     # blue accent
MUTED = RGBColor(0x5B, 0x61, 0x70)      # secondary text
LIGHT = RGBColor(0xF2, 0xF5, 0xFA)      # light fill for cards/bg
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GOOD = RGBColor(0x1E, 0x8E, 0x3E)       # green
BAD = RGBColor(0xC5, 0x3A, 0x2B)        # red

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf


def set_run(run, text, size, color=INK, bold=False, italic=False, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def accent_bar(slide):
    """Thin accent bar along the left edge of a content slide."""
    bar = slide.shapes.add_shape(1, 0, 0, Inches(0.18), SH)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    bar.shadow.inherit = False


def content_slide(title, kicker=None):
    slide = add_slide()
    bg(slide, WHITE)
    accent_bar(slide)
    # kicker (small label)
    if kicker:
        _, ktf = textbox(slide, Inches(0.7), Inches(0.45), Inches(11), Inches(0.4))
        kp = ktf.paragraphs[0]
        set_run(kp.add_run(), kicker.upper(), 13, ACCENT, bold=True)
    # title
    _, ttf = textbox(slide, Inches(0.65), Inches(0.8), Inches(12), Inches(1.1))
    tp = ttf.paragraphs[0]
    set_run(tp.add_run(), title, 34, INK, bold=True)
    return slide


def bullets(slide, items, top=2.1, left=0.85, width=11.6, size=20, gap=10):
    """items: list of (text, level, color, bold) or plain strings."""
    _, tf = textbox(slide, Inches(left), Inches(top), Inches(width), Inches(4.7))
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, str):
            text, level, color, bold = item, 0, INK, False
        else:
            text = item[0]
            level = item[1] if len(item) > 1 else 0
            color = item[2] if len(item) > 2 else INK
            bold = item[3] if len(item) > 3 else False
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(gap)
        bullet = "•  " if level == 0 else "–  "
        r = p.add_run()
        set_run(r, bullet + text, size - (level * 2), color, bold=bold)
    return tf


# ===========================================================================
# Slide 1 — Title
# ===========================================================================
s = add_slide()
bg(s, INK)
# accent block
block = s.shapes.add_shape(1, 0, Inches(2.55), SW, Inches(0.07))
block.fill.solid(); block.fill.fore_color.rgb = ACCENT; block.line.fill.background()
block.shadow.inherit = False

_, tf = textbox(s, Inches(0.9), Inches(2.7), Inches(11.5), Inches(2))
p = tf.paragraphs[0]
set_run(p.add_run(), "How the Spider Merge Bot Works", 46, WHITE, bold=True)
p2 = tf.add_paragraph()
p2.space_before = Pt(14)
set_run(p2.add_run(), "Automatic forward-merging across release branches", 24, RGBColor(0xC7, 0xD2, 0xF0))

_, tf3 = textbox(s, Inches(0.95), Inches(5.1), Inches(11), Inches(0.6))
set_run(tf3.paragraphs[0].add_run(), "A 10-minute intro for new developers", 18, RGBColor(0x9A, 0xA6, 0xC0), italic=True)

# ===========================================================================
# Slide 2 — The problem
# ===========================================================================
s = content_slide("The problem it solves", kicker="Why this exists")
bullets(s, [
    ("A fix often lands in an older release branch first — e.g. release-5.8.0.", 0),
    ("But that fix needs to flow forward to newer releases, and eventually to main.", 0),
    ("Doing this by hand across many branches is tedious, easy to forget, and easy to get wrong.", 0),
    ("The merge bot automates it — and gets a developer involved only when there's a real conflict.", 0, ACCENT, True),
], top=2.2, size=22, gap=16)

# ===========================================================================
# Slide 3 — What it does (happy path) + forward callout
# ===========================================================================
s = content_slide("What it does: the happy path", kicker="The core idea")
bullets(s, [
    ("When a PR merges, the bot automatically merges it forward through the release chain — all the way to main.", 0),
    ("No human needed for a clean merge. It just happens.", 0),
    ("If a conflict appears, it files an issue and hands you a clean workspace to resolve it.", 0),
], top=2.1, size=22, gap=16)

# "forward" callout card
card = s.shapes.add_shape(5, Inches(0.85), Inches(5.0), Inches(11.6), Inches(1.4))
card.fill.solid(); card.fill.fore_color.rgb = LIGHT
card.line.color.rgb = ACCENT; card.line.width = Pt(1.25)
card.shadow.inherit = False
ctf = card.text_frame; ctf.word_wrap = True
ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
cp = ctf.paragraphs[0]; cp.alignment = PP_ALIGN.LEFT
set_run(cp.add_run(), "Always forward:  ", 20, ACCENT, bold=True)
set_run(cp.add_run(), "newer branches receive the few new commits — never the other way around.", 20, INK)

# ===========================================================================
# Slide 4 — The golden rule
# ===========================================================================
s = content_slide("The golden rule for developers", kicker="Remember this one thing")

# big rule banner
banner = s.shapes.add_shape(5, Inches(0.85), Inches(1.9), Inches(11.6), Inches(1.5))
banner.fill.solid(); banner.fill.fore_color.rgb = INK
banner.line.fill.background(); banner.shadow.inherit = False
btf = banner.text_frame; btf.word_wrap = True; btf.vertical_anchor = MSO_ANCHOR.MIDDLE
bp = btf.paragraphs[0]; bp.alignment = PP_ALIGN.CENTER
set_run(bp.add_run(), "Always branch from  ", 24, WHITE, bold=True)
set_run(bp.add_run(), "branch-here-{version}", 24, RGBColor(0x7E, 0xB0, 0xFF), bold=True)
bp2 = btf.add_paragraph(); bp2.alignment = PP_ALIGN.CENTER
set_run(bp2.add_run(), "never from  ", 24, WHITE, bold=True)
set_run(bp2.add_run(), "release-{version}", 24, RGBColor(0xFF, 0x9B, 0x8A), bold=True)

bullets(s, [
    ("Why: branch-here only contains commits that made it all the way to main.", 0),
    ("So you never inherit anyone else's unresolved merge conflicts.", 0),
    ("main has no branch-here pointer — it's already the final destination, so branch from main directly.", 0, MUTED, False),
], top=3.7, size=20, gap=14)

# ===========================================================================
# Slide 5 — The branch zoo (table)
# ===========================================================================
s = content_slide("The branch zoo", kicker="Four kinds of branches")
rows = [
    ("Branch", "What it's for", "Who touches it"),
    ("release-*", "The actual releases & hotfixes", "Bot only"),
    ("branch-here-*", "Safe points for YOU to branch from", "You branch from it"),
    ("merge-forward-*", "An isolated merge chain per PR", "Bot only"),
    ("merge-conflicts-*", "A workspace to resolve a conflict", "You, when asked"),
]
table_shape = s.shapes.add_table(len(rows), 3, Inches(0.85), Inches(2.1),
                                 Inches(11.6), Inches(3.4))
table = table_shape.table
table.columns[0].width = Inches(3.5)
table.columns[1].width = Inches(5.3)
table.columns[2].width = Inches(2.8)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = table.cell(r, c)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
        para = cell.text_frame.paragraphs[0]
        run = para.add_run()
        is_header = (r == 0)
        mono = (c == 0 and not is_header)
        set_run(run, val, 16 if not is_header else 16,
                WHITE if is_header else INK,
                bold=is_header or mono,
                font="Consolas" if mono else "Calibri")
        cell.fill.solid()
        if is_header:
            cell.fill.fore_color.rgb = ACCENT
        else:
            cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT

_, ptf = textbox(s, Inches(0.85), Inches(5.9), Inches(11.6), Inches(0.8))
pp = ptf.paragraphs[0]
set_run(pp.add_run(), "Punchline:  ", 20, ACCENT, bold=True)
set_run(pp.add_run(), "you only ever touch branch-here-* (and a workspace if asked). The bot owns the rest.", 20, INK)

# ===========================================================================
# Slide 6 — When there's a conflict
# ===========================================================================
s = content_slide("When there's a conflict", kicker="Conflicts are normal, not failures")
bullets(s, [
    ("The bot hits a conflict it can't resolve automatically.", 0),
    ("It files an issue and creates a merge-conflicts workspace branch for you.", 0),
    ("You resolve the conflict in that workspace and open a PR.", 0),
    ("A later bot run picks up where it left off and completes the chain forward.", 0),
    ("A conflict is expected behavior — an issue was created, not an error.", 0, GOOD, True),
], top=2.1, size=22, gap=16)

# ===========================================================================
# Slide 7 — Under the hood + recovery
# ===========================================================================
s = content_slide("Under the hood & recovery", kicker="Just enough internals")
bullets(s, [
    ("It's a single GitHub Action that runs in two phases:", 0, INK, True),
    ("Phase 1 — Auto-merge: merge each PR forward through the release chain.", 1),
    ("Phase 2 — Branch maintenance: advance branch-here pointers, update releases, clean up.", 1),
    ("Configuration lives in .spider-merge-bot-config.json in the Spider Impact repo.", 0),
    ("If branches get out of sync, manual-merge.sh resets everything to a pristine state.", 0),
    ("It's idempotent — if a conflict occurs, resolve it, commit, and re-run.", 1, MUTED, False),
], top=2.1, size=20, gap=12)

# ===========================================================================
# Slide 8 — Cheat sheet
# ===========================================================================
s = add_slide()
bg(s, INK)
_, tf = textbox(s, Inches(0.9), Inches(0.7), Inches(11.5), Inches(1))
set_run(tf.paragraphs[0].add_run(), "Cheat sheet", 38, WHITE, bold=True)

cheats = [
    ("1", "Branch from branch-here-*", "Never from release-*. main is fine to branch from directly."),
    ("2", "Let the bot merge forward", "Clean merges flow to main automatically — no action needed."),
    ("3", "Resolve conflicts where it tells you", "The bot files an issue + a workspace branch. Fix there, open a PR."),
    ("4", "Use manual-merge.sh to recover", "Resets branches to a pristine state. Idempotent — safe to re-run."),
]
top = 1.9
for num, head, body in cheats:
    # number chip
    chip = s.shapes.add_shape(9, Inches(0.9), Inches(top), Inches(0.7), Inches(0.7))
    chip.fill.solid(); chip.fill.fore_color.rgb = ACCENT; chip.line.fill.background()
    chip.shadow.inherit = False
    ctf = chip.text_frame; ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
    cpp = ctf.paragraphs[0]; cpp.alignment = PP_ALIGN.CENTER
    set_run(cpp.add_run(), num, 22, WHITE, bold=True)
    # text
    _, ttf = textbox(s, Inches(1.85), Inches(top - 0.07), Inches(10.8), Inches(1.2))
    hp = ttf.paragraphs[0]
    set_run(hp.add_run(), head, 22, WHITE, bold=True)
    bp = ttf.add_paragraph()
    set_run(bp.add_run(), body, 16, RGBColor(0xB8, 0xC2, 0xDC))
    top += 1.25

out = "docs/spider-merge-bot.pptx"
prs.save(out)
print(f"Saved {out} with {len(prs.slides._sldIdLst)} slides")
